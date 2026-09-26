"""Build the writer's JavaScript rules from the reviewed Excel workbook."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


DEFAULT_INPUT = Path(__file__).resolve().parents[1] / "Rule_Review.xlsx"
DEFAULT_OUTPUT = Path(__file__).resolve().parents[1] / "static" / "rules.js"
SOURCE_SHEET = "Sheet1"
SEPARATOR = ";"


def split_values(value: Any) -> list[str]:
	"""Split a workbook cell into cleaned semicolon-separated values."""
	if value is None:
		return []
	return [part.strip() for part in str(value).split(SEPARATOR) if part.strip()]


def build_rules(input_path: Path) -> list[dict[str, Any]]:
	"""Read the rule rows from the workbook and convert them to frontend rules."""
	workbook = load_workbook(input_path, data_only=True, read_only=True)
	if SOURCE_SHEET not in workbook.sheetnames:
		raise ValueError(f"Workbook does not contain the {SOURCE_SHEET!r} sheet")

	worksheet = workbook[SOURCE_SHEET]
	rows = worksheet.iter_rows(values_only=True)
	headers = next((row for row in rows if any(value is not None for value in row)), None)
	if not headers:
		return []

	column_indexes = {str(header).strip().lower(): index for index, header in enumerate(headers) if header is not None}
	required_columns = {"word", "level", "reason", "replacement"}
	missing_columns = required_columns - column_indexes.keys()
	if missing_columns:
		missing = ", ".join(sorted(missing_columns))
		raise ValueError(f"Missing required columns: {missing}")

	rules = []
	for rule_number, row in enumerate(rows, start=1):
		patterns = split_values(row[column_indexes["word"]] if column_indexes["word"] < len(row) else None)
		if not patterns:
			continue

		level = row[column_indexes["level"]] if column_indexes["level"] < len(row) else None
		reason = row[column_indexes["reason"]] if column_indexes["reason"] < len(row) else None
		replacements = split_values(row[column_indexes["replacement"]] if column_indexes["replacement"] < len(row) else None)
		severity = str(level).strip().lower() if level is not None and str(level).strip() else "strong"
		if severity not in {"strong", "hint"}:
			severity = "strong"

		rules.append(
			{
				"id": f"R{rule_number:02d}",
				"patterns": patterns,
				"replacements": replacements,
				"severity": severity,
				"reason": "" if reason is None else str(reason).strip(),
			}
		)

	return rules


def write_rules(output_path: Path, rules: list[dict[str, Any]]) -> None:
	"""Write rules as a JavaScript module consumed by the writer frontend."""
	output_path.parent.mkdir(parents=True, exist_ok=True)
	serialized_rules = json.dumps(rules, ensure_ascii=False, indent=2)
	output_path.write_text(
		"(function () {\n"
		'  "use strict";\n'
		"  const D = (window.Decadre = window.Decadre || {});\n"
		f"  D.rules = {serialized_rules};\n"
		"})();\n",
		encoding="utf-8",
	)


def main() -> None:
	parser = argparse.ArgumentParser(description=__doc__)
	parser.add_argument("input", nargs="?", type=Path, default=DEFAULT_INPUT, help="Input Rule_Review.xlsx path")
	parser.add_argument("output", nargs="?", type=Path, default=DEFAULT_OUTPUT, help="Output rules.js path")
	args = parser.parse_args()

	rules = build_rules(args.input)
	write_rules(args.output, rules)
	print(f"Wrote {len(rules)} rules to {args.output}")


if __name__ == "__main__":
	main()
